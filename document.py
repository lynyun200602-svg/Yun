"""按「简历集合」分子目录存储 pdf/docx/图片，读取指定集合的全部文档文本作为上下文。

目录结构：data/upload_files/<collection_id>/<文件名>
图片通过 OCR（RapidOCR，自带中英文模型）转成文字后参与问答，识别结果缓存为同名 .txt。
"""
import os
import shutil

from pypdf import PdfReader
from docx import Document

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
UPLOAD_ROOT = os.path.join(BASE_DIR, "data", "upload_files")

IMG_EXTS = {".jpg", ".jpeg", ".png", ".webp", ".bmp"}
SUPPORTED = {".pdf", ".docx"} | IMG_EXTS


def upload_dir(collection_id=None):
    """返回集合的文件目录；collection_id 为 None 时返回根目录。"""
    if collection_id is None:
        return UPLOAD_ROOT
    d = os.path.join(UPLOAD_ROOT, str(collection_id))
    os.makedirs(d, exist_ok=True)
    return d


def list_files(collection_id):
    d = upload_dir(collection_id)
    files = []
    if not os.path.isdir(d):
        return files
    for name in sorted(os.listdir(d)):
        ext = os.path.splitext(name)[1].lower()
        if ext in SUPPORTED:
            p = os.path.join(d, name)
            if os.path.isfile(p):
                files.append({"name": name, "size": os.path.getsize(p)})
    return files


def _extract_pdf(path):
    reader = PdfReader(path)
    parts = []
    for page in reader.pages:
        t = page.extract_text() or ""
        parts.append(t)
    return "\n".join(parts)


def _extract_docx(path):
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


_ocr_engine = None


def _get_ocr():
    """懒加载 OCR 引擎（首次识别图片时才初始化，加载较慢）。"""
    global _ocr_engine
    if _ocr_engine is None:
        from rapidocr_onnxruntime import RapidOCR
        _ocr_engine = RapidOCR()
    return _ocr_engine


def _extract_image(path):
    """用 OCR 提取图片文字，结果缓存到同名 .txt，避免重复识别。"""
    cache_path = path + ".txt"
    if os.path.exists(cache_path) and os.path.getmtime(cache_path) >= os.path.getmtime(path):
        with open(cache_path, "r", encoding="utf-8") as f:
            return f.read()

    engine = _get_ocr()
    result, _ = engine(path)
    text = "\n".join(item[1] for item in result) if result else ""

    try:
        with open(cache_path, "w", encoding="utf-8") as f:
            f.write(text)
    except Exception:
        pass
    return text


def read_documents(collection_id):
    """返回指定集合拼接后的全部文档文本；没有文档时返回 None。"""
    chunks = []
    for f in list_files(collection_id):
        p = os.path.join(upload_dir(collection_id), f["name"])
        ext = os.path.splitext(f["name"])[1].lower()
        try:
            if ext == ".pdf":
                text = _extract_pdf(p)
            elif ext == ".docx":
                text = _extract_docx(p)
            else:
                text = _extract_image(p)
        except Exception as e:  # 单个文件解析失败不阻断整体
            chunks.append(f"[文件 {f['name']} 解析失败：{e}]")
            continue
        chunks.append(f"===== 文档：{f['name']} =====\n{text}")
    if not chunks:
        return None
    return "\n\n".join(chunks)


def delete_collection_files(collection_id):
    d = upload_dir(collection_id)
    if os.path.isdir(d):
        shutil.rmtree(d)


def migrate_legacy_files(default_collection_id):
    """把旧版平铺在根目录的文件迁入默认集合子目录。"""
    root = UPLOAD_ROOT
    if not os.path.isdir(root):
        return 0
    moved = 0
    for name in os.listdir(root):
        p = os.path.join(root, name)
        if os.path.isfile(p) and os.path.splitext(name)[1].lower() in SUPPORTED:
            dest = os.path.join(upload_dir(default_collection_id), name)
            shutil.move(p, dest)
            moved += 1
    return moved
